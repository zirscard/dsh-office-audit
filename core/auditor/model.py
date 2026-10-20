"""Workbook model: sheet metadata, header/period detection, label indexing."""
from __future__ import annotations

import re

from openpyxl import load_workbook

from .util import (
    DAY_WORDS,
    MULTIPLE_WORDS,
    MONEY_WORDS,
    STRONG_MONEY_WORDS,
    RATIO_WORDS,
    is_formula,
    is_number,
    normalize_label,
    period_of_header,
    to_number,
)

HEADER_SCAN_ROWS = 15


class SheetData:
    """One worksheet pair (formulas + cached values)."""

    def __init__(self, ws, wsv):
        self.ws = ws
        self.wsv = wsv
        self.name = ws.title
        self.max_row = ws.max_row or 0
        self.max_column = ws.max_column or 0
        self.header_row = self._detect_header_row()
        self.period_columns = self._detect_periods()
        self._labels = None
        self._merged = list(ws.merged_cells.ranges) if ws.merged_cells else []

    # -- cells ---------------------------------------------------------
    def formula_cell(self, row: int, col: int):
        return self.ws.cell(row=row, column=col)

    def value(self, row: int, col: int):
        """Cached (recalculated) value when available, else the formula-cell value."""
        if self.wsv is not None:
            v = self.wsv.cell(row=row, column=col).value
            if v is not None:
                return v
        return self.ws.cell(row=row, column=col).value

    def number(self, row: int, col: int):
        return to_number(self.value(row, col))

    def is_formula(self, row: int, col: int) -> bool:
        cell = self.ws.cell(row=row, column=col)
        return cell.data_type == "f" or is_formula(cell.value)

    def formula_text(self, row: int, col: int):
        v = self.ws.cell(row=row, column=col).value
        return v if is_formula(v) else None

    def number_format(self, row: int, col: int):
        return self.ws.cell(row=row, column=col).number_format

    # -- layout --------------------------------------------------------
    def _detect_header_row(self):
        best_row, best_score = None, 0
        limit = min(self.max_row, HEADER_SCAN_ROWS)
        for row in range(1, limit + 1):
            score = 0
            for col in range(1, self.max_column + 1):
                v = self.ws.cell(row=row, column=col).value
                if isinstance(v, str) and len(v.strip()) > 0 and not is_formula(v):
                    if re.search(r"[A-Za-z\u4e00-\u9fff]", v):
                        score += 1
            if score > best_score:
                best_row, best_score = row, score
        return best_row if best_score >= 2 else None

    def _detect_periods(self):
        periods = {}
        if self.header_row is None:
            return periods
        for col in range(1, self.max_column + 1):
            text = self.ws.cell(row=self.header_row, column=col).value
            period = period_of_header(text)
            if period is not None:
                periods[col] = period
        return periods

    # -- labels --------------------------------------------------------
    def _merged_lookup(self, row, col):
        for rng in self._merged:
            if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
                return self.ws.cell(row=rng.min_row, column=rng.min_col)
        return None

    def label(self, row: int) -> str:
        """First textual cell in columns 1..3 (merged-aware, follows empty merged anchors)."""
        for col in range(1, min(self.max_column, 3) + 1):
            v = self.ws.cell(row=row, column=col).value
            if isinstance(v, str) and v.strip() and not is_formula(v):
                return v.strip()
            if v is None:
                anchor = self._merged_lookup(row, col)
                if anchor is not None and isinstance(anchor.value, str) and anchor.value.strip():
                    return anchor.value.strip()
        return ""

    def labels(self):
        if self._labels is None:
            self._labels = [(row, self.label(row)) for row in range(1, self.max_row + 1)]
        return self._labels

    def label_rows(self):
        """[(row, label)] for every non-empty row label."""
        return [(row, label) for row, label in self.labels() if label]

    def find_label_row(self, words, *, exclude=(), strict=False) -> int | None:
        """Find the first row whose label matches.

        strict=True requires the normalized label to START with one of the
        words — needed for totals, so 「流动资产合计」 does not match the
        「资产合计」 pattern used for 「资产总计」.
        """
        for row, label in self.label_rows():
            norm = normalize_label(label)
            if any(bad in norm for bad in exclude):
                continue
            if strict:
                if any(norm.startswith(word) for word in words):
                    return row
            elif any(word in norm for word in words):
                return row
        return None

    def has_data(self) -> bool:
        return self.max_row > 1 and self.max_column > 1

    # -- content classification ---------------------------------------
    def row_kind(self, row: int) -> str:
        """Classify a row label into ratio/day/multiple/money/other."""
        label = self.label(row)
        if not label:
            return "other"
        norm = normalize_label(label)
        if any(w in norm for w in DAY_WORDS):
            return "days"
        if any(w in norm for w in ("校验", "检查")):
            return "ratio"
        if any(w in norm for w in MULTIPLE_WORDS):
            return "multiple"
        if (
            any(w in norm for w in RATIO_WORDS)
            or ("率" in norm and "汇率" not in norm)
            or norm.endswith(("度", "率", "比", "倍数", "天数", "乘数"))
            or "/" in norm
            or "÷" in norm
            or any(w in norm for w in ("增速", "增长", "同比", "环比"))
        ):
            return "ratio"
        if any(w in norm for w in MONEY_WORDS):
            return "money"
        return "other"

    def numeric_columns(self):
        """Columns with at least 3 numeric cells below the header."""
        counts = {}
        start = (self.header_row or 0) + 1
        for row in range(start, self.max_row + 1):
            for col in self.period_columns or range(2, self.max_column + 1):
                if is_number(self.value(row, col)):
                    counts[col] = counts.get(col, 0) + 1
        return {col for col, n in counts.items() if n >= 3}


class WorkbookModel:
    def __init__(self, path, recalc_path=None, limits=None):
        self.path = path
        self.limits = limits or {"max_cells": 400_000, "max_formula_scan": 120_000}
        self.wb = load_workbook(path, data_only=False, read_only=False)
        self.wbv = load_workbook(recalc_path, data_only=True, read_only=False) if recalc_path else None
        self.sheet_names = list(self.wb.sheetnames)
        self.sheets = {}
        for name in self.sheet_names:
            ws = self.wb[name]
            wsv = self.wbv[name] if (self.wbv is not None and name in self.wbv.sheetnames) else None
            self.sheets[name] = SheetData(ws, wsv)
        self._defined_names = self._collect_defined_names()

    def _collect_defined_names(self):
        names = {}
        try:
            for name, defined in self.wb.defined_names.items():
                text = getattr(defined, "attr_text", None) or getattr(defined, "value", None)
                if text:
                    names[name] = str(text)
        except Exception:  # noqa: BLE001 - defined-name API varies across openpyxl versions
            pass
        return names

    def sheet(self, name: str) -> SheetData:
        return self.sheets[name]

    def iter_sheets(self):
        for name in self.sheet_names:
            yield self.sheets[name]

    def all_formula_cells(self):
        scanned = 0
        for sheet in self.iter_sheets():
            for row in range(1, sheet.max_row + 1):
                for col in range(1, sheet.max_column + 1):
                    cell = sheet.formula_cell(row, col)
                    if cell.data_type == "f" or is_formula(cell.value):
                        yield sheet, row, col, cell
                        scanned += 1
                        if scanned >= self.limits["max_formula_scan"]:
                            return

    def count_formulas(self):
        n = 0
        for sheet in self.iter_sheets():
            for row in range(1, sheet.max_row + 1):
                for col in range(1, sheet.max_column + 1):
                    if sheet.formula_cell(row, col).data_type == "f":
                        n += 1
        return n
