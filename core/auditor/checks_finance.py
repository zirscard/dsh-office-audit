"""Three-statement and ratio-table checks for financial models."""
from __future__ import annotations

import re

from openpyxl.formula import Tokenizer

from .checks_generic import cell_ref
from .util import (
    LABELS,
    is_derived_sheet_name,
    format_compact,
    is_number,
    normalize_label,
    to_number,
)

YEAR_DAY_LITERALS = {"365", "360", "366"}
RATIO_ROW_KINDS = {"ratio", "days", "multiple"}


STRICT_TOTAL_KEYS = {"assets_total", "liabilities_total", "equity_total", "liab_equity_total"}


def _find(sheet, key):
    return sheet.find_label_row(LABELS[key], strict=key in STRICT_TOTAL_KEYS)


def _ordered_periods(sheet, *, kind=None):
    items = []
    for col, period in sheet.period_columns.items():
        if kind is not None and period[0] != kind:
            continue
        items.append((col, period))
    rank = {"year": 0, "quarter": 1, "month": 2}.get
    items.sort(key=lambda item: (item[1][1], rank(item[1][0], 0), item[1][2] if len(item[1]) > 2 else 0))
    return items


def _period_label(period) -> str:
    kind, year = period[0], period[1]
    if kind == "quarter" and len(period) > 2:
        return f"{year}Q{period[2]}"
    if kind == "month" and len(period) > 2:
        return f"{year}-{period[2]:02d}"
    return str(year)


def _value_at(sheet, row, period):
    """Find (value, col) for a period on a sheet; exact match first, then kind/year."""
    if sheet is None:
        return None, None
    for col, candidate in sheet.period_columns.items():
        if candidate == period:
            value = sheet.number(row, col)
            if value is not None:
                return value, col
    for col, candidate in sheet.period_columns.items():
        if candidate[0] == period[0] and candidate[1] == period[1]:
            value = sheet.number(row, col)
            if value is not None:
                return value, col
    return None, None


def _columns_with_numbers(sheet, row):
    cols = []
    for col in range(2, sheet.max_column + 1):
        if to_number(sheet.value(row, col)) is not None:
            cols.append(col)
    return cols


def check_balance_sheet(model, result, opts):
    if model.wbv is None:
        return
    for sheet in model.iter_sheets():
        if is_derived_sheet_name(sheet.name):
            continue
        assets_row = _find(sheet, "assets_total")
        if assets_row is None:
            continue
        liabilities_row = _find(sheet, "liabilities_total")
        equity_row = _find(sheet, "equity_total")
        combined_row = _find(sheet, "liab_equity_total")
        has_three = liabilities_row is not None and equity_row is not None
        if not has_three and combined_row is None:
            continue

        periods = _ordered_periods(sheet)
        if not periods:
            periods = [(col, ("col", col)) for col in _columns_with_numbers(sheet, assets_row)]
        mismatches = []
        for col, period in periods:
            assets = to_number(sheet.value(assets_row, col))
            if assets is None:
                continue
            if has_three:
                liabilities = to_number(sheet.value(liabilities_row, col))
                equity = to_number(sheet.value(equity_row, col))
                if liabilities is None or equity is None:
                    continue
                diff = assets - (liabilities + equity)
                bound = max(liabilities + equity, 1.0)
            else:
                combined = to_number(sheet.value(combined_row, col))
                if combined is None:
                    continue
                diff = assets - combined
                bound = max(abs(assets), 1.0)
            tolerance = max(1.0, abs(bound) * 0.001)
            if abs(diff) > tolerance:
                label = sheet.period_columns.get(col, ("", _period_label(period)))[1]
                mismatches.append((col, label, diff, assets, tolerance))
        if mismatches:
            col, label, diff, assets, tolerance = mismatches[0]
            evidence = "；".join(
                f"{label}：资产 {format_compact(a)}，差异 {format_compact(d)}（容差 {format_compact(t)}）"
                for _c, label, d, a, t in mismatches[:6]
            )
            if len(mismatches) > 6:
                evidence += f"；…共 {len(mismatches)} 期不平衡"
            result.add(
                "critical",
                "balance-sheet",
                f"资产负债表不平衡（{len(mismatches)} 期）",
                sheet=sheet.name,
                cell=cell_ref(assets_row, col),
                detail="资产总计 ≠ 负债合计 + 所有者权益合计，差额超过 0.1% 容差；这是三表模型最基础的硬伤。",
                evidence=evidence,
                suggestion="先核对差额是否等于某笔漏结转（如当期利润、分红、汇兑差），再检查合计行公式是否漏行/重复行。",
            )
        else:
            result.add(
                "info",
                "balance-sheet",
                "资产负债表平衡（或未识别到不平衡）",
                sheet=sheet.name,
                detail=f"已检查 {len(periods)} 期，资产与负债+权益在 0.1% 容差内一致或数据不足。",
            )


def check_retained_earnings(model, result, opts):
    if model.wbv is None:
        return
    for sheet in model.iter_sheets():
        if is_derived_sheet_name(sheet.name):
            continue
        re_row = _find(sheet, "retained_earnings")
        if re_row is None:
            continue
        ni_row = _find(sheet, "net_income")
        ni_sheet = sheet
        if ni_row is None:
            for other in model.iter_sheets():
                candidate = _find(other, "net_income")
                if candidate is not None:
                    ni_row, ni_sheet = candidate, other
                    break
        div_row = _find(sheet, "dividends")
        div_sheet = sheet
        if div_row is None:
            for other in model.iter_sheets():
                candidate = _find(other, "dividends")
                if candidate is not None:
                    div_row, div_sheet = candidate, other
                    break

        periods = _ordered_periods(sheet)
        if len(periods) < 2:
            continue
        mismatches = []
        for idx in range(1, len(periods)):
            prev_col, prev_period = periods[idx - 1]
            cur_col, cur_period = periods[idx]
            re_prev = to_number(sheet.value(re_row, prev_col))
            re_cur = to_number(sheet.value(re_row, cur_col))
            if re_prev is None or re_cur is None:
                continue
            ni = to_number(ni_sheet.value(ni_row, cur_col)) if ni_row is not None and ni_sheet is sheet else None
            if ni is None and ni_row is not None:
                ni, _ = _value_at(ni_sheet, ni_row, cur_period)
            if ni is None:
                continue
            div = 0.0
            if div_row is not None:
                div_value = to_number(div_sheet.value(div_row, cur_col)) if div_sheet is sheet else None
                if div_value is None:
                    div_value, _ = _value_at(div_sheet, div_row, cur_period)
                if div_value is not None:
                    div = div_value
            expected = re_prev + ni - div
            diff = re_cur - expected
            tolerance = max(1.0, abs(re_cur) * 0.001, abs(expected) * 0.001)
            if abs(diff) > tolerance:
                mismatches.append((cur_col, _period_label(cur_period), re_cur, expected, diff))
        if mismatches:
            col, label, actual, expected, diff = mismatches[0]
            evidence = "；".join(
                f"{lbl}：期末 {format_compact(actual)}，应为 {format_compact(exp)}，差 {format_compact(d)}"
                for _c, lbl, actual, exp, d in mismatches[:5]
            )
            if len(mismatches) > 5:
                evidence += f"；…共 {len(mismatches)} 期"
            result.add(
                "warning",
                "retained-earnings",
                f"留存收益滚动不闭合（{len(mismatches)} 期）",
                sheet=sheet.name,
                cell=cell_ref(re_row, col),
                detail="期末留存收益 ≠ 期初留存收益 + 当期净利润 − 分红。常见原因：法定口径 vs 模拟/FDD 口径混用、漏提盈余公积、期初数被手改。",
                evidence=evidence,
                suggestion="确认口径（法定 or 模拟）并把全表统一到同一口径；若为 FDD 调整口径，需单独列调节表，不能直接混入法定表。",
            )


def check_cash_tieout(model, result, opts):
    if model.wbv is None:
        return
    for sheet in model.iter_sheets():
        if is_derived_sheet_name(sheet.name):
            continue
        bs_cash_row = _find(sheet, "cash_bs")
        if bs_cash_row is None:
            continue
        cf_end_row, cf_sheet = None, None
        for candidate_sheet in model.iter_sheets():
            row = _find(candidate_sheet, "cash_cf_end")
            if row is not None:
                cf_end_row, cf_sheet = row, candidate_sheet
                break
        if cf_end_row is None:
            continue
        periods = _ordered_periods(sheet)
        mismatches = []
        for col, period in periods:
            bs_cash = to_number(sheet.value(bs_cash_row, col))
            if bs_cash is None:
                continue
            cf_cash, cf_col = _value_at(cf_sheet, cf_end_row, period)
            if cf_cash is None:
                continue
            diff = bs_cash - cf_cash
            tolerance = max(1.0, abs(cf_cash) * 0.001)
            if abs(diff) > tolerance:
                mismatches.append((col, _period_label(period), bs_cash, cf_cash, diff))
        if mismatches:
            col, label, bs_cash, cf_cash, diff = mismatches[0]
            evidence = "；".join(
                f"{lbl}：BS {format_compact(b)} vs CF {format_compact(c)}，差 {format_compact(d)}"
                for _c, lbl, b, c, d in mismatches[:6]
            )
            if len(mismatches) > 6:
                evidence += f"；…共 {len(mismatches)} 期"
            result.add(
                "critical",
                "cash-tieout",
                f"现金勾稽不符（{len(mismatches)} 期）",
                sheet=sheet.name,
                cell=cell_ref(bs_cash_row, col),
                detail=f"资产负债表「货币资金/现金」与现金流量表「期末现金及现金等价物余额」对不上（工作表 {sheet.name} vs {cf_sheet.name}）。",
                evidence=evidence,
                suggestion="检查受限资金/定期存款是否被排除在现金等价物外、CF 期末现金公式是否漏加期初，以及是否有未入 CF 的现金变动。",
            )


def check_cash_flow_subtotals(model, result, opts):
    if model.wbv is None:
        return
    for sheet in model.iter_sheets():
        if is_derived_sheet_name(sheet.name):
            continue
        op_row = _find(sheet, "cf_operating")
        inv_row = _find(sheet, "cf_investing")
        fin_row = _find(sheet, "cf_financing")
        net_row = _find(sheet, "cash_cf_net")
        periods = _ordered_periods(sheet)
        if not periods:
            continue

        if op_row and inv_row and fin_row and net_row:
            mismatches = []
            for col, period in periods:
                values = [to_number(sheet.value(row, col)) for row in (op_row, inv_row, fin_row, net_row)]
                if any(v is None for v in values):
                    continue
                op, inv, fin, net = values
                diff = net - (op + inv + fin)
                if abs(diff) > max(1.0, abs(net) * 0.001):
                    mismatches.append((col, _period_label(period), net, op + inv + fin, diff))
            if mismatches:
                col, label, net, total, diff = mismatches[0]
                result.add(
                    "critical",
                    "cash-flow-subtotals",
                    f"现金流量表净增加额不等于三大活动之和（{len(mismatches)} 期）",
                    sheet=sheet.name,
                    cell=cell_ref(net_row, col),
                    detail="现金及现金等价物净增加额 ≠ 经营 + 投资 + 筹资净额。",
                    evidence="；".join(
                        f"{lbl}：净增加 {format_compact(n)}，三项合计 {format_compact(t)}，差 {format_compact(d)}"
                        for _c, lbl, n, t, d in mismatches[:6]
                    ),
                    suggestion="检查净增加额是否为手填常数、三大活动小计是否漏行或重复计算。",
                )

        begin_row = _find(sheet, "cash_cf_begin")
        end_row = _find(sheet, "cash_cf_end")
        net_row = _find(sheet, "cash_cf_net")
        if begin_row and end_row and net_row:
            mismatches = []
            for col, period in periods:
                begin = to_number(sheet.value(begin_row, col))
                net = to_number(sheet.value(net_row, col))
                end = to_number(sheet.value(end_row, col))
                if begin is None or net is None or end is None:
                    continue
                diff = end - (begin + net)
                if abs(diff) > max(1.0, abs(end) * 0.001):
                    mismatches.append((col, _period_label(period), end, begin + net, diff))
            if mismatches:
                col, label, end, expected, diff = mismatches[0]
                result.add(
                    "critical",
                    "cash-flow-subtotals",
                    f"期末现金 ≠ 期初现金 + 净增加额（{len(mismatches)} 期）",
                    sheet=sheet.name,
                    cell=cell_ref(end_row, col),
                    detail="现金流量表期末现金的纵向勾稽不闭合。",
                    evidence="；".join(
                        f"{lbl}：期末 {format_compact(e)}，期初+净增 {format_compact(x)}，差 {format_compact(d)}"
                        for _c, lbl, e, x, d in mismatches[:6]
                    ),
                    suggestion="确认期初现金是否等于上期期末、净增加额是否引用正确小计。",
                )


def check_ratio_conventions(model, result, opts):
    revenue_threshold = opts.get("revenue_threshold", 1_000_000)
    for sheet in model.iter_sheets():
        ratio_rows = [
            (row, sheet.label(row))
            for row in range(1, sheet.max_row + 1)
            if sheet.row_kind(row) in RATIO_ROW_KINDS and sheet.label(row)
        ]
        if not ratio_rows:
            continue
        periods = _ordered_periods(sheet)
        kinds = {period[0] for _col, period in periods}
        if "year" in kinds and ("quarter" in kinds or "month" in kinds):
            result.add(
                "info",
                "ratio-conventions",
                "同一表混用年度与季度/月度期间",
                sheet=sheet.name,
                detail="年度列与季度列相邻时，周转率/收益率若用季度值直接替代年度值会放大或缩小 4 倍。",
                suggestion="在比率表增加「期间类型」参数：季度值先年化（×4 或 ×12/月数），或统一只用年度列。",
            )

        # Hard-coded year-day literals inside ratio formulas.
        hardcoded = []
        for row, label in ratio_rows:
            for col in range(2, sheet.max_column + 1):
                formula = sheet.formula_text(row, col)
                if formula is None:
                    continue
                try:
                    tokens = Tokenizer(formula).items
                except Exception:  # noqa: BLE001
                    continue
                for token in tokens:
                    if token.type == "OPERAND" and token.subtype == "NUMBER" and token.value.strip() in YEAR_DAY_LITERALS:
                        hardcoded.append((row, col, token.value, formula))
                        break
        if hardcoded:
            row, col, literal, formula = hardcoded[0]
            result.add(
                "warning",
                "ratio-conventions",
                f"比率公式内联硬编码年天数 {literal}（{len(hardcoded)} 处）",
                sheet=sheet.name,
                cell=cell_ref(row, col),
                detail="年天数直接写进公式，改口径（365→360）要逐格改，容易漏改。",
                evidence="；".join(
                    f"{cell_ref(r, c)} ={lit}" for r, c, lit, _f in hardcoded[:6]
                ),
                suggestion="新建「参数」表集中年天数（365/360）、年化倍数、阈值，公式改为引用参数单元格。",
            )

        # Mixed average-balance vs period-end convention.
        average_like = []
        end_like = []
        for row, _label in ratio_rows:
            for col in range(2, sheet.max_column + 1):
                formula = sheet.formula_text(row, col)
                if formula is None:
                    continue
                upper = formula.upper()
                if "AVERAGE" in upper or re.search(r"\([^()]+\)\s*/\s*2\b", upper):
                    average_like.append((row, col, formula))
                elif re.search(r"/[A-Z$]+\d+", upper):
                    end_like.append((row, col, formula))
        if average_like and end_like:
            r, c, f = average_like[0]
            result.add(
                "info",
                "ratio-conventions",
                "平均余额与期末余额口径混用",
                sheet=sheet.name,
                cell=cell_ref(r, c),
                detail=f"有 {len(average_like)} 个公式用平均余额（AVERAGE 或 /2），{len(end_like)} 个直接用期末余额；同一张比率表口径应统一。",
                evidence=f"平均口径示例 {cell_ref(r, c)} {f}",
                suggestion="在参数表定义「余额口径」切换项，周转类统一用平均余额，收益率类统一用期末或平均，并在表头注明。",
            )

        # Pre-revenue / tiny-denominator periods: only flag ratios that are
        # genuinely meaningless (unit-agnostic: compare revenue against its own
        # median, and ratios against their own scale).
        revenue_row = _find(sheet, "revenue")
        revenue_sheet = sheet
        if revenue_row is None:
            for other in model.iter_sheets():
                candidate = _find(other, "revenue")
                if candidate is not None:
                    revenue_row, revenue_sheet = candidate, other
                    break
        if revenue_row is not None and periods and model.wbv is not None:
            series = []
            for col, period in periods:
                if revenue_sheet is sheet:
                    revenue = to_number(sheet.value(revenue_row, col))
                else:
                    revenue, _rc = _value_at(revenue_sheet, revenue_row, period)
                if revenue is not None and revenue > 0:
                    series.append((col, period, revenue))
            if len(series) >= 2:
                positive = sorted(item[2] for item in series)
                median = positive[len(positive) // 2]
                tiny = {col: (period, revenue) for col, period, revenue in series if revenue < median * 0.15}
                absurd = []
                for row, label in ratio_rows:
                    values = [to_number(sheet.value(row, col)) for col, _p in periods]
                    finite = [abs(v) for v in values if v is not None]
                    scale = max(finite) if finite else 0.0
                    for col, _p in periods:
                        if col not in tiny:
                            continue
                        value = to_number(sheet.value(row, col))
                        if value is None:
                            continue
                        kind = sheet.row_kind(row)
                        is_absurd = (
                            (kind == "days" and abs(value) > 3650)
                            or (kind == "multiple" and abs(value) > 100)
                            or (kind == "ratio" and abs(value) > 10)
                            or (scale > 0 and abs(value) > scale * 5)
                        )
                        if is_absurd:
                            absurd.append((row, col, label, value, tiny[col][1], tiny[col][0]))
                if absurd:
                    row, col, label, value, revenue, period = absurd[0]
                    result.add(
                        "warning",
                        "ratio-conventions",
                        f"未投产/微收入期的比率失真（{len(absurd)} 处）",
                        sheet=sheet.name,
                        cell=cell_ref(row, col),
                        detail=f"该期营业收入仅 {format_compact(revenue)}（不足中位数的 15%），分母过小导致比率被放大（如 {label} = {format_compact(value)}）；不应作为趋势或决策依据。",
                        evidence="；".join(
                            f"{cell_ref(r, c)} {lbl}={format_compact(v)}（营收 {format_compact(rev)}，{_period_label(p)}）"
                            for r, c, lbl, v, rev, p in absurd[:5]
                        ),
                        suggestion="对未投产期用 IF/IFERROR 或条件格式标注「不适用」，并在参数表登记未投产收入阈值；趋势图起点从投产期开始。",
                    )

        # Ratio cells whose formula exists but the period has no value (often /0).
        empty_cells = []
        intentional_texts = ("需核查", "不适用（未投产期）", "不适用", "N/A", "-", "—")
        if model.wbv is None:
            ratio_rows = []
        for row, _label in ratio_rows:
            for col, _period in periods:
                if sheet.formula_text(row, col) is None:
                    continue
                raw = sheet.value(row, col)
                if isinstance(raw, str) and raw.strip() in intentional_texts:
                    continue
                if to_number(raw) is None:
                    empty_cells.append((row, col))
        if empty_cells:
            row, col = empty_cells[0]
            result.add(
                "info",
                "ratio-conventions",
                f"比率行存在无法计算的期间（{len(empty_cells)} 处）",
                sheet=sheet.name,
                cell=cell_ref(row, col),
                detail="这些单元格有公式但没有可用结果，通常是分母为 0/空或首期缺数；直接显示会留下 #DIV/0! 或空白。",
                evidence="、" .join(cell_ref(r, c) for r, c in empty_cells[:8]),
                suggestion="用 IFERROR 或判空逻辑显示「不适用」，并在表下注明首期/未投产期不计算比率。",
            )
