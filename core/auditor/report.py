"""Report rendering: markdown + the in-workbook 模型审计 sheet."""
from __future__ import annotations

import datetime as _dt

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .findings import CATEGORY_LABELS, SEVERITIES

SEVERITY_TITLES = {"critical": "严重（Critical）", "warning": "警告（Warning）", "info": "提示（Info）"}
SEVERITY_FILLS = {
    "critical": "FFC7CE",
    "warning": "FFEB9C",
    "info": "DDEBF7",
}


def render_markdown(result, *, engine=None) -> str:
    summary = result.summary()
    lines = []
    lines.append(f"# 工作簿审计报告：{result.workbook.rsplit('/', 1)[-1]}")
    lines.append("")
    lines.append(f"- 工作簿：`{result.workbook}`")
    lines.append(f"- 审计模式：{result.profile}")
    lines.append(f"- 生成时间：{_dt.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    if engine:
        lines.append(f"- 引擎：{engine}")
    recalc = result.meta.get("recalculated")
    lines.append(f"- 公式重算：{'是（LibreOffice Kit）' if recalc else '否（静态检查）'}")
    lines.append("")
    lines.append("## 汇总")
    lines.append("")
    lines.append("| 级别 | 数量 |")
    lines.append("| --- | ---: |")
    for sev in SEVERITIES:
        lines.append(f"| {SEVERITY_TITLES[sev]} | {summary[sev]} |")
    lines.append(f"| **合计** | **{summary['total']}** |")
    lines.append("")

    if summary["total"] == 0:
        lines.append("未发现需要报告的问题。")
        lines.append("")
    else:
        for sev in SEVERITIES:
            group = [f for f in result.sorted_findings() if f.severity == sev]
            if not group:
                continue
            lines.append(f"## {SEVERITY_TITLES[sev]}")
            lines.append("")
            for finding in group:
                where = " · ".join(
                    part for part in (finding.sheet, finding.cell) if part
                )
                lines.append(f"### {finding.id} {finding.title}")
                lines.append("")
                meta_bits = [f"类别：{CATEGORY_LABELS.get(finding.category, finding.category)}"]
                if where:
                    meta_bits.append(f"位置：`{where}`")
                lines.append(" | ".join(meta_bits))
                if finding.detail:
                    lines.append("")
                    lines.append(finding.detail)
                if finding.evidence:
                    lines.append("")
                    lines.append(f"- 证据：{finding.evidence}")
                if finding.suggestion:
                    lines.append(f"- 建议：{finding.suggestion}")
                lines.append("")

    lines.append("## 工作表概览")
    lines.append("")
    lines.append("| 工作表 | 行 | 列 | 公式 | 常数 |")
    lines.append("| --- | ---: | ---: | ---: | ---: |")
    for sheet in result.sheets:
        lines.append(
            f"| {sheet.get('name', '')} | {sheet.get('rows', 0)} | {sheet.get('columns', 0)} | "
            f"{sheet.get('formulaCells', 0)} | {sheet.get('constantCells', 0)} |"
        )
    lines.append("")
    lines.append("---")
    lines.append("")
    lines.append("由 dsh-office-audit 生成。审计不修改原文件；修复后请重跑复验。")
    lines.append("")
    return "\n".join(lines)


def write_audit_sheet(result, source_path: str, target_path: str):
    """Add/replace a 「模型审计」 sheet in a COPY of the workbook."""
    keep_vba = source_path.lower().endswith(".xlsm")
    wb = load_workbook(source_path, keep_vba=keep_vba)
    sheet_name = "模型审计"
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name, 0)

    headers = ["序号", "级别", "类别", "工作表", "单元格", "问题", "详情", "证据", "建议"]
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for finding in result.sorted_findings():
        ws.append([
            finding.id,
            {"critical": "严重", "warning": "警告", "info": "提示"}.get(finding.severity, finding.severity),
            CATEGORY_LABELS.get(finding.category, finding.category),
            finding.sheet or "",
            finding.cell or "",
            finding.title,
            finding.detail,
            finding.evidence,
            finding.suggestion,
        ])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        level = row[1].value
        fill_rgb = {"严重": SEVERITY_FILLS["critical"], "警告": SEVERITY_FILLS["warning"], "提示": SEVERITY_FILLS["info"]}.get(level)
        for idx, cell in enumerate(row):
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
            if idx == 1 and fill_rgb:
                cell.fill = PatternFill("solid", fgColor=fill_rgb)
                cell.font = Font(bold=True)
    widths = [8, 8, 18, 14, 10, 34, 52, 56, 44]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, ws.max_row)}"

    # A small header block above? No: keep the sheet machine-readable from row 1.
    wb.save(target_path)
    return target_path
